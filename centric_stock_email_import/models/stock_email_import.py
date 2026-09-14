import base64
import csv
import hashlib
import io
import logging
import re
import zipfile
from collections import OrderedDict

from markupsafe import Markup

from odoo import _, api, fields, models
from odoo.exceptions import UserError, ValidationError
from odoo.tools import float_compare, float_is_zero

from .stock_email_import_column import TARGETS, normalize_header

_logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = (".xlsx", ".xlsm", ".csv")
HEADER_SCAN_ROWS = 20
PRODUCT_KEYS = ("product_code", "barcode", "product_name")


def _cell_text(value):
    """Excel stores codes like 10045 as 10045.0 - keep them as '10045'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _cell_number(value):
    """Return a float, None for an empty cell, or raise ValueError."""
    if isinstance(value, bool):
        raise ValueError(value)
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"[^\d,.\-]", "", str(value))
    if not text:
        if str(value).strip():
            raise ValueError(value)
        return None
    if "," in text and "." not in text:
        # 1,250 (thousands) versus 12,5 (decimal comma)
        if re.fullmatch(r"-?\d{1,3}(,\d{3})+", text):
            text = text.replace(",", "")
        else:
            text = text.replace(",", ".")
    else:
        text = text.replace(",", "")
    return float(text)


class StockEmailImport(models.Model):
    """One received stocktake sheet and what happened to it."""

    _name = "stock.email.import"
    _description = "Stock Email Import"
    _inherit = ["mail.thread"]
    _order = "id desc"

    name = fields.Char(required=True, default=lambda self: _("Stocktake"))
    config_id = fields.Many2one(
        "stock.email.import.config", required=True, ondelete="restrict",
        default=lambda self: self.env["stock.email.import.config"].search(
            [("company_id", "=", self.env.company.id)], limit=1),
    )
    company_id = fields.Many2one(
        "res.company", related="config_id.company_id", store=True, readonly=True,
    )
    currency_id = fields.Many2one(related="company_id.currency_id")
    email_from = fields.Char(string="Sender", readonly=True)
    received_date = fields.Datetime(default=fields.Datetime.now, readonly=True)
    file = fields.Binary(string="Sheet", attachment=True)
    file_name = fields.Char()
    checksum = fields.Char(compute="_compute_checksum", store=True, index=True)
    allow_duplicate = fields.Boolean(
        help="Apply this sheet even though the same file was already applied.",
    )
    state = fields.Selection(
        [
            ("draft", "Draft"),
            ("pending", "Waiting"),
            ("done", "Applied"),
            ("failed", "Failed"),
            ("rejected", "Rejected"),
        ],
        default="draft", required=True, tracking=True, index=True,
    )
    error_message = fields.Text(readonly=True)
    applied_date = fields.Datetime(readonly=True)
    line_ids = fields.One2many("stock.email.import.line", "import_id")
    move_ids = fields.Many2many("stock.move", string="Inventory Moves", readonly=True)
    line_count = fields.Integer(compute="_compute_totals")
    error_count = fields.Integer(compute="_compute_totals")
    changed_count = fields.Integer(compute="_compute_totals")
    value_change = fields.Monetary(
        compute="_compute_totals",
        help="Estimated at product cost: sum of (counted - previous) x cost.",
    )

    @api.depends("file")
    def _compute_checksum(self):
        for record in self:
            record.checksum = (
                hashlib.sha256(base64.b64decode(record.file)).hexdigest()
                if record.file else False
            )

    @api.depends("line_ids.status", "line_ids.difference_qty", "line_ids.value_change")
    def _compute_totals(self):
        for record in self:
            lines = record.line_ids
            record.line_count = len(lines)
            record.error_count = len(lines.filtered(lambda l: l.status == "error"))
            record.changed_count = len(lines.filtered(lambda l: l.difference_qty))
            record.value_change = sum(lines.mapped("value_change"))

    # ------------------------------------------------------------------
    # Mail gateway
    # ------------------------------------------------------------------

    @api.model
    def _is_supported_file(self, filename):
        return (filename or "").lower().endswith(SUPPORTED_EXTENSIONS)

    @api.model
    def message_new(self, msg_dict, custom_values=None):
        """Turn an email into imports: one per supported attachment.

        Nothing is applied here. The mail gateway runs inside the fetchmail
        transaction, so a slow or failing sheet must not bounce the email;
        the records are queued and the processing cron is woken instead.
        """
        values = dict(custom_values or {})
        Config = self.env["stock.email.import.config"].sudo()
        config = Config.browse(values.get("config_id")).exists()
        if not config:
            config = Config.search([("company_id", "=", self.env.company.id)], limit=1)
        values["config_id"] = config.id
        values.pop("company_id", None)  # related to the config

        email_from = msg_dict.get("email_from") or ""
        files = [
            (attachment[0], attachment[1])
            for attachment in msg_dict.get("attachments") or []
            if self._is_supported_file(attachment[0])
        ]
        values.update({
            "name": msg_dict.get("subject") or _("Emailed stocktake"),
            "email_from": email_from,
            "state": "pending",
        })
        if not config or not config._is_sender_allowed(email_from):
            values.update({
                "state": "rejected",
                "error_message": _(
                    "%(sender)s is not an allowed sender for stocktake emails.",
                    sender=email_from or _("Unknown sender"),
                ),
            })
            files = files[:1]
        elif not files:
            values.update({
                "state": "failed",
                "error_message": _(
                    "The email has no Excel (.xlsx) or CSV attachment. Old .xls "
                    "files must be saved as .xlsx first."
                ),
            })

        extra_files = []
        if files:
            first, extra_files = files[0], files[1:]
            values.update(self._file_values(*first))

        record = super(StockEmailImport, self.sudo()).message_new(
            msg_dict, custom_values=values)
        for filename, content in extra_files:
            record.copy(dict(self._file_values(filename, content),
                             name=values["name"], email_from=email_from,
                             state="pending", received_date=record.received_date))
        if values["state"] == "pending":
            self.env.ref(
                "centric_stock_email_import.ir_cron_process_stock_email_imports"
            ).sudo()._trigger()
        return record

    @api.model
    def _file_values(self, filename, content):
        if isinstance(content, str):
            content = content.encode("utf-8")
        return {"file": base64.b64encode(content), "file_name": filename}

    # ------------------------------------------------------------------
    # Processing
    # ------------------------------------------------------------------

    @api.model
    def _cron_process_pending(self, limit=20):
        self.search([("state", "=", "pending")], order="id", limit=limit)._process()

    def action_process(self):
        for record in self:
            if record.state not in ("draft", "failed", "pending"):
                raise UserError(_("Only draft, waiting or failed imports can be applied."))
            if not record.file:
                raise UserError(_("Attach a sheet first."))
        self._process()
        return True

    def action_reset_to_draft(self):
        self.filtered(lambda r: r.state in ("failed", "rejected")).write({
            "state": "draft", "error_message": False,
        })

    def action_view_moves(self):
        self.ensure_one()
        return {
            "type": "ir.actions.act_window",
            "name": _("Inventory Moves"),
            "res_model": "stock.move",
            "view_mode": "list,form",
            "domain": [("id", "in", self.move_ids.ids)],
        }

    def _process(self):
        for record in self:
            record = record.sudo()
            config = record.config_id
            record.line_ids.unlink()
            record.error_message = False
            try:
                importer = record.with_user(config.user_id).with_company(record.company_id)
                importer._run_import()
            except (UserError, ValidationError) as error:
                record._mark_failed(str(error.args[0] if error.args else error))
            except Exception as error:  # noqa: BLE001 - reported on the record
                _logger.exception("Stock email import %s failed", record.id)
                record._mark_failed(_("Unexpected error: %(error)s", error=error))

    def _mark_failed(self, message):
        self.ensure_one()
        self.write({"state": "failed", "error_message": message})
        self._report(_("Stocktake not applied"), message)

    def _run_import(self):
        """Validate every row, then apply all of them or none of them."""
        self.ensure_one()
        self._check_duplicate()
        rows = self._validate_rows(self._read_sheet())
        errors = [row for row in rows if row["status"] == "error"]
        if errors:
            self.env["stock.email.import.line"].sudo().create(
                [self._line_values(row) for row in rows])
            raise UserError(_(
                "%(count)s row(s) have errors, so no stock was changed. "
                "Fix the rows listed on the import and send the sheet again.",
                count=len(errors),
            ))
        if not any(row["status"] == "ok" for row in rows):
            raise UserError(_("The sheet has no rows with a counted quantity."))

        with self.env.cr.savepoint():
            self._apply(rows)
        self.env["stock.email.import.line"].sudo().create(
            [self._line_values(row) for row in rows])
        self.sudo().write({"state": "done", "applied_date": fields.Datetime.now()})
        self._report(_("Stocktake applied"), self._summary_text())

    def _check_duplicate(self):
        if self.allow_duplicate or not self.checksum:
            return
        previous = self.search([
            ("id", "!=", self.id),
            ("checksum", "=", self.checksum),
            ("company_id", "=", self.company_id.id),
            ("state", "=", "done"),
        ], limit=1)
        if previous:
            raise UserError(_(
                "This exact file was already applied by '%(name)s' on %(date)s. "
                "Re-applying an old count would undo the stock movements since. "
                "Tick 'Allow Duplicate' on the import to apply it anyway.",
                name=previous.name, date=previous.applied_date,
            ))

    # -- reading -------------------------------------------------------

    def _read_sheet(self):
        """Return (header row index, {target: column index}, rows)."""
        raw = base64.b64decode(self.file or b"")
        filename = (self.file_name or "").lower()
        if filename.endswith(".csv"):
            sheets = [self._read_csv(raw)]
        else:
            sheets = self._read_xlsx(raw)

        header_map = self.env["stock.email.import.column"].sudo()._header_map()
        for rows in sheets:
            for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
                columns = {}
                for column, cell in enumerate(row):
                    target = header_map.get(normalize_header(cell))
                    if target and target not in columns:
                        columns[target] = column
                if "quantity" in columns and columns.keys() & set(PRODUCT_KEYS):
                    return index, columns, rows
        raise UserError(_(
            "No header row found. The sheet needs a quantity column and a "
            "product code, barcode or name column. Accepted headers are set in "
            "Inventory > Stock Email Import > Column Mapping."
        ))

    def _read_xlsx(self, raw):
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError):
            raise UserError(_(
                "'%(file)s' is not a readable .xlsx file. Old .xls files must "
                "be saved as .xlsx first.", file=self.file_name,
            ))
        try:
            return [
                [list(row) for row in sheet.iter_rows(values_only=True)]
                for sheet in workbook.worksheets
            ]
        finally:
            workbook.close()

    def _read_csv(self, raw):
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [row for row in csv.reader(io.StringIO(text), dialect)]

    # -- validating ----------------------------------------------------

    def _validate_rows(self, sheet):
        header_index, columns, rows = sheet
        config = self.config_id
        caches = {"product": {}, "location": {}, "lot": {}}
        result = []
        for index in range(header_index + 1, len(rows)):
            raw = rows[index]

            def get(target, raw=raw):
                column = columns.get(target)
                return raw[column] if column is not None and column < len(raw) else None

            row = {target: _cell_text(get(target)) for target, _label in TARGETS
                   if target not in ("quantity", "price")}
            quantity_cell, price_cell = get("quantity"), get("price")
            if not any(row[key] for key in PRODUCT_KEYS) and _cell_text(quantity_cell) == "":
                continue  # blank row

            row.update({
                "location_text": row.pop("location"),
                "lot_text": row.pop("lot"),
                "row_number": index + 1,
                "status": "ok",
                "messages": [],
                "product": self.env["product.product"],
                "new_product": False,
                "location": self.env["stock.location"],
                "lot": self.env["stock.lot"],
                "new_lot": False,
                "counted_qty": 0.0,
                "price": False,
            })
            result.append(row)

            self._validate_quantity(row, quantity_cell)
            self._validate_price(row, price_cell, config)
            self._resolve_product(row, config, caches["product"])
            self._resolve_location(row, config, caches["location"])
            if row["product"] or row["new_product"]:
                self._resolve_lot(row, config, caches["lot"])

            if any(message[0] == "error" for message in row["messages"]):
                row["status"] = "error"
            elif row["status"] != "skipped":
                row["status"] = "ok"
        self._check_serial_totals(result)
        return result

    @staticmethod
    def _add(row, level, message):
        row["messages"].append((level, message))

    def _validate_quantity(self, row, cell):
        try:
            quantity = _cell_number(cell)
        except ValueError:
            self._add(row, "error", _("Quantity '%(value)s' is not a number.", value=cell))
            return
        if quantity is None:
            row["status"] = "skipped"
            self._add(row, "info", _("No quantity: row ignored, stock unchanged."))
        elif quantity < 0:
            self._add(row, "error", _("A counted quantity cannot be negative."))
        else:
            row["counted_qty"] = quantity

    def _validate_price(self, row, cell, config):
        if config.price_update == "none":
            return
        try:
            price = _cell_number(cell)
        except ValueError:
            self._add(row, "error", _("Price '%(value)s' is not a number.", value=cell))
            return
        if price is not None and price < 0:
            self._add(row, "error", _("A price cannot be negative."))
        elif price is not None:
            row["price"] = price

    def _resolve_product(self, row, config, cache):
        key = (row["product_code"], row["barcode"], row["product_name"])
        if key not in cache:
            cache[key] = self._find_product(row, config)
        product, level, message = cache[key]
        if product == "new":
            row["new_product"] = key
        else:
            row["product"] = product
        if message:
            self._add(row, level, message)

    def _find_product(self, row, config):
        """Return (product | 'new', message level, message)."""
        Product = self.env["product.product"]
        company_domain = [("company_id", "in", [False, config.company_id.id])]
        product = Product
        if row["product_code"]:
            product = Product.search(
                [("default_code", "=", row["product_code"])] + company_domain, limit=2)
        if not product and row["barcode"]:
            product = Product.search(
                [("barcode", "=", row["barcode"])] + company_domain, limit=2)
        if not product and row["product_name"] and not (row["product_code"] or row["barcode"]):
            product = Product.search(
                [("name", "=", row["product_name"])] + company_domain, limit=2)

        label = row["product_code"] or row["barcode"] or row["product_name"]
        if len(product) > 1:
            return Product, "error", _(
                "'%(product)s' matches more than one product.", product=label)
        if product:
            if not product.is_storable:
                return Product, "error", _(
                    "%(product)s does not track inventory (not a storable product).",
                    product=product.display_name)
            return product, None, None
        if config.create_missing_products and row["product_name"]:
            return "new", "info", _("New product will be created.")
        return Product, "error", _("Product '%(product)s' not found.", product=label)

    def _resolve_location(self, row, config, cache):
        text = row["location_text"]
        if text not in cache:
            cache[text] = self._find_location(text, config)
        location, message = cache[text]
        row["location"] = location
        if message:
            self._add(row, "error", message)

    def _find_location(self, text, config):
        Location = self.env["stock.location"]
        if not text:
            location = config.sudo()._default_location()
            if not location:
                return Location, _("No default location: set one on the configuration.")
            return self.env["stock.location"].browse(location.id), None
        locations = Location.search([
            ("usage", "=", "internal"),
            ("company_id", "=", config.company_id.id),
            "|", "|",
            ("complete_name", "=", text),
            ("barcode", "=", text),
            ("name", "=", text),
        ])
        if len(locations) > 1:
            exact = locations.filtered(lambda l: l.complete_name == text)
            locations = exact if len(exact) == 1 else locations
        if len(locations) > 1:
            return Location, _(
                "Location '%(location)s' is ambiguous; use its full name, e.g. %(example)s.",
                location=text, example=locations[0].complete_name)
        if not locations:
            return Location, _("Location '%(location)s' not found.", location=text)
        return locations, None

    def _resolve_lot(self, row, config, cache):
        product = row["product"]
        tracking = product.tracking if product else "none"
        lot_name = row["lot_text"]
        if tracking == "none":
            if lot_name:
                self._add(row, "info", _(
                    "Lot '%(lot)s' ignored: product is not tracked by lot.", lot=lot_name))
            return
        if not lot_name:
            self._add(row, "error", _(
                "%(product)s is tracked by lot/serial: a lot number is required.",
                product=product.display_name))
            return
        if tracking == "serial" and float_compare(row["counted_qty"], 1, 2) > 0:
            self._add(row, "error", _("A serial number can only be counted once."))
        key = (product.id, lot_name)
        if key not in cache:
            cache[key] = self.env["stock.lot"].search([
                ("name", "=", lot_name),
                ("product_id", "=", product.id),
                ("company_id", "in", [False, config.company_id.id]),
            ], limit=1)
        if cache[key]:
            row["lot"] = cache[key]
        elif config.create_missing_lots:
            row["new_lot"] = lot_name
            self._add(row, "info", _("New lot '%(lot)s' will be created.", lot=lot_name))
        else:
            self._add(row, "error", _("Lot '%(lot)s' not found.", lot=lot_name))

    def _check_serial_totals(self, rows):
        seen = {}
        for row in rows:
            product = row["product"]
            if row["status"] != "ok" or not product or product.tracking != "serial":
                continue
            key = (product.id, row["lot"].id or row["new_lot"])
            if key in seen and row["counted_qty"] and seen[key]:
                row["status"] = "error"
                self._add(row, "error", _(
                    "Serial number also counted on row %(row)s.", row=seen[key]))
            elif row["counted_qty"]:
                seen[key] = row["row_number"]

    # -- applying ------------------------------------------------------

    def _apply(self, rows):
        config = self.config_id
        ok_rows = [row for row in rows if row["status"] == "ok"]
        self._create_missing_products(ok_rows, config)
        self._update_prices(ok_rows, config)
        self._create_missing_lots(ok_rows, config)

        groups = OrderedDict()
        for row in ok_rows:
            key = (row["product"].id, row["location"].id, row["lot"].id)
            groups.setdefault(key, []).append(row)

        Quant = self.env["stock.quant"].with_context(inventory_mode=True)
        Move = self.env["stock.move"]
        last_move = Move.sudo().search([], order="id desc", limit=1)
        digits = self.env["decimal.precision"].precision_get("Product Unit")
        quants = Quant.browse()
        for (product_id, location_id, lot_id), group in groups.items():
            product = group[0]["product"]
            counted = sum(row["counted_qty"] for row in group)
            previous = sum(Quant.search([
                ("product_id", "=", product_id),
                ("location_id", "=", location_id),
                ("lot_id", "=", lot_id or False),
                ("package_id", "=", False),
                ("owner_id", "=", False),
            ]).mapped("quantity"))
            quant = Quant.create({
                "product_id": product_id,
                "location_id": location_id,
                "lot_id": lot_id or False,
                "inventory_quantity": counted,
            })
            quant.inventory_quantity_set = True
            quants |= quant

            difference = counted - previous
            if float_is_zero(difference, precision_digits=digits):
                difference = 0.0
            first = group[0]
            first.update({
                "previous_qty": previous,
                "difference_qty": difference,
                "value_change": difference * product.standard_price,
            })
            if len(group) > 1:
                rows_text = ", ".join(str(row["row_number"]) for row in group)
                for row in group:
                    self._add(row, "info", _(
                        "Rows %(rows)s are the same product, location and lot; "
                        "their quantities are added together (total %(total)s).",
                        rows=rows_text, total=counted))

        result = quants.with_context(
            inventory_name=_("Stock Email Import: %(name)s", name=self.name),
        ).action_apply_inventory()
        if isinstance(result, dict):
            raise UserError(_(
                "Odoo asked for a confirmation (%(action)s) before applying the "
                "count, so nothing was applied.", action=result.get("name") or result.get("res_model"),
            ))

        moves = Move.sudo().search([
            ("id", ">", last_move.id or 0),
            ("product_id", "in", [key[0] for key in groups]),
            "|",
            ("location_id.usage", "=", "inventory"),
            ("location_dest_id.usage", "=", "inventory"),
        ])
        self.sudo().move_ids = [fields.Command.set(moves.ids)]

    def _create_missing_products(self, rows, config):
        new_products = {}
        for row in rows:
            key = row["new_product"]
            if not key:
                continue
            if key not in new_products:
                values = {
                    "name": row["product_name"],
                    "default_code": row["product_code"] or False,
                    "barcode": row["barcode"] or False,
                    "type": "consu",
                    "is_storable": True,
                    "company_id": False,
                }
                if row["category"]:
                    values["categ_id"] = self._find_or_create_category(row["category"]).id
                new_products[key] = self.env["product.product"].create(values)
            row["product"] = new_products[key]
            row["new_product"] = False

    def _find_or_create_category(self, name):
        Category = self.env["product.category"]
        category = Category.search([("complete_name", "=", name)], limit=1) \
            or Category.search([("name", "=", name)], limit=1)
        return category or Category.create({"name": name})

    def _update_prices(self, rows, config):
        if config.price_update == "none":
            return
        for row in rows:
            if row["price"] is False:
                continue
            product = row["product"]
            if config.price_update == "cost":
                if float_compare(product.standard_price, row["price"], 6):
                    product.standard_price = row["price"]
            elif float_compare(product.list_price, row["price"], 6):
                product.product_tmpl_id.list_price = row["price"]

    def _create_missing_lots(self, rows, config):
        created = {}
        for row in rows:
            if not row["new_lot"]:
                continue
            key = (row["product"].id, row["new_lot"])
            if key not in created:
                created[key] = self.env["stock.lot"].create({
                    "name": row["new_lot"],
                    "product_id": row["product"].id,
                    "company_id": config.company_id.id,
                })
            row["lot"] = created[key]
            row["new_lot"] = False

    # -- reporting -----------------------------------------------------

    def _line_values(self, row):
        levels = [level for level, _message in row["messages"]]
        return {
            "import_id": self.id,
            "row_number": row["row_number"],
            "product_code": row["product_code"],
            "barcode": row["barcode"],
            "product_name": row["product_name"],
            "category_name": row["category"],
            "location_name": row["location"].complete_name or row["location_text"] or False,
            "lot_name": row["lot"].name or row["new_lot"] or row["lot_text"] or False,
            "product_id": row["product"].id if row["product"] else False,
            "location_id": row["location"].id if row["location"] else False,
            "lot_id": row["lot"].id if row["lot"] else False,
            "counted_qty": row["counted_qty"],
            "previous_qty": row.get("previous_qty", 0.0),
            "difference_qty": row.get("difference_qty", 0.0),
            "value_change": row.get("value_change", 0.0),
            "price": row["price"] or 0.0,
            "status": "error" if "error" in levels else row["status"],
            "message": "\n".join(message for _level, message in row["messages"]),
        }

    def _summary_text(self):
        self.invalidate_recordset(["line_ids"])
        return _(
            "%(lines)s row(s) read, %(changed)s quantity change(s), estimated "
            "value change %(value)s %(currency)s.",
            lines=self.line_count, changed=self.changed_count,
            value=round(self.value_change, 2), currency=self.currency_id.name or "",
        )

    def _report(self, title, text):
        """Log the outcome on the chatter and, if enabled, email the sender."""
        self.ensure_one()
        record = self.sudo()
        record.invalidate_recordset(["line_ids"])
        body = Markup("<p><strong>%s</strong></p><p>%s</p>") % (title, text)
        errors = record.line_ids.filtered(lambda l: l.status == "error")
        if errors:
            items = Markup("").join(
                Markup("<li>%s %s: %s</li>") % (
                    _("Row"), line.row_number, line.message)
                for line in errors[:50]
            )
            body += Markup("<ul>%s</ul>") % items
            if len(errors) > 50:
                body += Markup("<p>%s</p>") % _("... and %(count)s more.", count=len(errors) - 50)
        record.message_post(body=body)

        config = record.config_id
        if not (config.notify_sender and record.email_from and record.state in ("done", "failed")):
            return
        self.env["mail.mail"].sudo().create({
            "subject": "%s: %s" % (title, record.name),
            "body_html": body,
            "email_to": record.email_from,
            "email_from": config.company_id.email_formatted or config.user_id.email_formatted,
            "auto_delete": True,
        })
